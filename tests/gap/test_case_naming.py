from pathlib import Path
import json
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from src.GaP.libs.case_naming import CaseFiles, case_files, resolve_case_stems, sanitize_name
from src.GaP.libs.models.simulation_scenario import SimulationDesign


def test_sanitize_name_is_filesystem_safe():
    assert sanitize_name("NO 32/4-1") == "NO_32_4_1"
    assert sanitize_name("  hot case!! ") == "hot_case"
    assert sanitize_name("///") == "case"


def test_stems_keep_readable_names_and_index_only_collisions():
    stems = resolve_case_stems("Smeaheia", ["baseline", "hot case", "hot-case", "Baseline 2", "HOT_CASE"])

    assert stems == ["Smeaheia_baseline", "Smeaheia_hot_case_02", "Smeaheia_hot_case_03", "Smeaheia_Baseline_2", "Smeaheia_HOT_CASE_05"]
    assert resolve_case_stems("Smeaheia", ["baseline", "hot case", "hot-case", "Baseline 2", "HOT_CASE"]) == stems


def test_design_rejects_duplicate_case_names():
    with pytest.raises(ValueError, match="duplicated: baseline"):
        SimulationDesign(scenarios=[{"case_name": "baseline"}, {"case_name": "baseline"}])


def test_case_files_fall_back_to_template_names(tmp_path):
    assert case_files(tmp_path) == CaseFiles.template()
    (tmp_path / "scenario.json").write_text(json.dumps({"case_metadata": {"case_stem": "well_a"}}), encoding="utf-8")
    assert case_files(tmp_path).prefix == Path("model/well_a")


def _fake_cirrus(tmp_path: Path, fixture_prefix: Path) -> Path:
    runner = tmp_path / "fake-cirrus"
    runner.write_text(
        "#!/bin/sh\n"
        f'cp "{fixture_prefix}.EGRID" "${{1%.in}}.EGRID"\n'
        f'cp "{fixture_prefix}.INIT" "${{1%.in}}.INIT"\n',
        encoding="utf-8",
    )
    runner.chmod(0o755)
    return runner


def _ten_case_workbook(root: Path, workbook: Path) -> list[str]:
    shutil.copyfile(root / "test_data/examples/wildcat/wildcat_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    sheet = excel["DesignMatrix"]
    columns = {sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)}
    template_row = [sheet.cell(2, column).value for column in range(1, sheet.max_column + 1)]
    sheet.delete_rows(2, sheet.max_row)
    names = [f"sens {index}" for index in range(8)] + ["hot case", "hot-case"]
    for name in names:
        row = list(template_row)
        row[columns["case_name"] - 1] = name
        sheet.append(row)
    excel.save(workbook)
    return names


def test_batch_creates_unique_navigable_cases_and_is_deterministic(tmp_path):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "ten_cases.xlsx"
    names = _ten_case_workbook(root, workbook)
    runner = _fake_cirrus(tmp_path, root / "test_data/examples/wildcat/model/TEMP-0")
    output_root = tmp_path / "batch"
    command = [
        sys.executable,
        "runscripts/run_workbook_scenarios_batch.py",
        "--xlsx",
        str(workbook),
        "--output-root",
        str(output_root),
        "--sim-command",
        f"{runner} {{deck}}",
        "--jobs",
        "5",
        "--force",
    ]

    subprocess.run(command, cwd=root, check=True, capture_output=True, text=True)
    manifest = json.loads((output_root / "batch_manifest.json").read_text(encoding="utf-8"))
    case_dirs = sorted(path.name for path in output_root.iterdir() if path.is_dir())

    assert [case["case_name"] for case in manifest["cases"]] == names
    assert [case["case_index"] for case in manifest["cases"]] == list(range(1, 11))
    assert {case["status"] for case in manifest["cases"]} == {"succeeded"}
    assert sorted(case["case_dir"] for case in manifest["cases"]) == case_dirs
    assert len(case_dirs) == 10
    assert manifest["cases"][-2]["case_stem"] == "wildcat_hot_case_09"
    assert manifest["cases"][-1]["case_stem"] == "wildcat_hot_case_10"
    for case in manifest["cases"]:
        case_root = output_root / case["case_dir"]
        metadata = json.loads((case_root / "scenario.json").read_text(encoding="utf-8"))["case_metadata"]
        assert metadata["case_index"] == case["case_index"]
        assert metadata["case_name"] == case["case_name"]
        assert metadata["template_files"]["deck"] == "model/TEMP-0.in"
        assert (case_root / "model" / f"{case['case_stem']}.in").is_file()
        assert (case_root / "model" / f"{case['case_stem']}.EGRID").is_file()
        assert (case_root / "include" / f"{case['case_stem']}_LGR.grdecl").is_file()
        assert not list(case_root.rglob("TEMP*"))

    subprocess.run(command, cwd=root, check=True, capture_output=True, text=True)
    assert json.loads((output_root / "batch_manifest.json").read_text(encoding="utf-8")) == manifest
    assert sorted(path.name for path in output_root.iterdir() if path.is_dir()) == case_dirs

    validation = subprocess.run(
        [sys.executable, "runscripts/validate_scenario_outputs.py", "--output-root", str(output_root)],
        cwd=root,
        capture_output=True,
        text=True,
    )
    assert validation.returncode == 0, validation.stdout
