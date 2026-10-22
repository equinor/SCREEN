import importlib.util
from pathlib import Path
import re
import shutil
import subprocess
import sys

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

from src.WellClass.libs.utils.xlsx_parser import xlsx_grid_policy, xlsx_to_simulation_design, xlsx_to_well_model
from src.WellClass.libs.well_pressure.pressure_table import PressureTable
from src.GaP.libs.pressure_inputs import resolve_pressure_inputs


PREPARE_PATH = Path(__file__).parents[2] / "runscripts/prepare_init_case.py"
PREPARE_SPEC = importlib.util.spec_from_file_location("prepare_init_case", PREPARE_PATH)
prepare_init_case = importlib.util.module_from_spec(PREPARE_SPEC)
PREPARE_SPEC.loader.exec_module(prepare_init_case)


def _write_minimal_workbook(path: Path, *, salinity: float = 0.032) -> None:
    metadata = pd.DataFrame({"key": ["namespace", "name", "author"], "value": ["screen", "xlsx-test", "pytest"]})
    header = pd.DataFrame(
        {
            "key": [
                "unique_wellbore_identifier",
                "depth_reference_rkb",
                "depth_reference_rkb_unit",
                "ground_elevation",
                "ground_elevation_unit",
                "total_depth_rkb",
                "total_depth_rkb_unit",
            ],
            "value": ["NO 00/0-0", 27, "m", 105, "m", 3997, "m"],
        }
    )
    grid_policy = pd.DataFrame(
        {
            "key": [
                "top_depth",
                "target_dz_water",
                "target_dz_overburden",
                "target_dz_reservoir",
                "reservoir_permx",
                "overburden_permx",
                "cells_per_layer",
            ],
            "value": [4.0, 50.0, 60.0, 10.0, 1000.0, 0.001, 400],
        }
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        metadata.to_excel(writer, sheet_name="Metadata", index=False)
        header.to_excel(writer, sheet_name="Header", index=False)
        grid_policy.to_excel(writer, sheet_name="GridPolicy", index=False)
        pd.DataFrame(
            {
                "temperature_gradient": [31.0],
                "ground_temperature": [4.0],
                "salinity": [salinity],
                "salinity_basis": ["mass_fraction"],
                "z_fluid_contact": [2400.0],
                "p_fluid_contact": [210.0],
                "overburden_datum_depth": [500.0],
                "z_resrv": [2400.0],
                "p_resrv": [210.0],
            }
        ).to_excel(writer, sheet_name="SubsurfaceAssumptions", index=False)
        pd.DataFrame(
            {
                "name": ["RESERVOIR"],
                "top_rkb": [1027.0],
                "bottom_rkb": [1400.0],
                "unit_type": ["reservoir"],
            }
        ).to_excel(writer, sheet_name="Stratigraphy", index=False)
        pd.DataFrame(
            {
                "name": ["Hole"],
                "type": ["hole"],
                "top_rkb": [444.0],
                "bottom_rkb": [1812.0],
                "diameter_in": [17.5],
                "shoe": [False],
            }
        ).to_excel(writer, sheet_name="HoleCasings", index=False)


def test_prepare_init_case_from_xlsx_stages_files(tmp_path):
    workbook = tmp_path / "well_input.xlsx"
    _write_minimal_workbook(workbook)

    output_root = tmp_path / "staged_case"
    repo_root = Path(__file__).parents[2]

    command = [
        sys.executable,
        "runscripts/prepare_init_case_from_xlsx.py",
        "--xlsx",
        str(workbook),
        "--output-root",
        str(output_root),
        "--write-well-json",
    ]
    subprocess.run(command, check=True, cwd=repo_root, capture_output=True, text=True)

    deck = output_root / "model" / "xlsx_test_default.in"
    grdecl = output_root / "include" / "xlsx_test_default_GRD.grdecl"
    co2_database = output_root / "include" / "co2_db_new.dat"
    tops = output_root / "include" / "tops_dz.inc"
    well_json = output_root / "well_input.json"

    assert deck.exists()
    assert grdecl.exists()
    assert co2_database.exists()
    assert tops.exists()
    assert well_json.exists()
    assert not list(output_root.rglob("TEMP*"))

    recipe = tops.read_text(encoding="utf-8")
    assert "TOPS 4" in recipe
    assert "1200*33.6667" in recipe
    assert "6000*59.6667" in recipe
    assert "16000*10" in recipe
    deck_text = deck.read_text(encoding="utf-8")
    assert "DATABASE ../include/co2_db_new.dat" in deck_text
    assert "\n BRINE  0.032 MASS\n" in deck_text
    assert "TYPE grdecl ../include/xlsx_test_default_GRD.grdecl" in deck_text
    assert deck_text.startswith("# SCREEN case: xlsx_test_default | scenario 'default' (DesignMatrix row 1) | template ")
    assert "_LGR.grdecl" not in grdecl.read_text(encoding="utf-8")
    grdecl_text = grdecl.read_text(encoding="utf-8")
    assert "EQLNUM 1 1 20 1 20 1 18 /" in grdecl_text
    assert "EQLNUM 2 1 20 1 20 19 58 /" in grdecl_text
    assert "PERMX 10000 1 20 1 20 1 3 /" in grdecl_text
    assert "PERMX 0.001 1 20 1 20 4 18 /" in grdecl_text
    assert "PORO 1 1 20 1 20 1 3 /" in grdecl_text
    assert "FIPLEG 3 1 20 1 20 19 58 /" in grdecl_text
    assert "FIPLEG 5 1 20 1 20 58 58 /" in grdecl_text
    assert "FIPLEG 7 1 1 1 20 1 18 /" in grdecl_text
    assert "PORV 2000 1 20 1 9 58 58 /" in grdecl_text  # the bottom layer, centre cell left out
    assert "PORV 2000 1 20 1 20 58 58 /" not in grdecl_text
    assert "PERMX 1000 1 20 1 20 56 58 /" in grdecl_text
    assert "TRANZ 0 1 20 1 20 19 19 /" in grdecl_text
    assert "PERMZ 0.1 1 20 1 20 4 58 /" in grdecl_text


@pytest.mark.parametrize("values", [
    {"z_fluid_contact": 1400.0},
    {"z_fluid_contact": 1400.0, "p_fluid_contact": 125.0},
    {"z_fluid_contact": 1400.0, "z_resrv": 1217.0, "p_resrv": 125.0},
    {"z_resrv": 1217.0, "p_resrv": 125.0},
])
def test_standalone_deck_uses_resolved_contact_pressure(tmp_path, values):
    workbook = tmp_path / "well.xlsx"
    _write_minimal_workbook(workbook)
    excel = load_workbook(workbook)
    sheet = excel["SubsurfaceAssumptions"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for field in ("z_fluid_contact", "p_fluid_contact", "z_resrv", "p_resrv"):
        sheet.cell(2, columns[field]).value = values.get(field)
    excel.save(workbook)
    result = subprocess.run(
        [sys.executable, "runscripts/prepare_init_case_from_xlsx.py", "--xlsx", str(workbook),
         "--output-root", str(tmp_path / "case")],
        cwd=Path(__file__).parents[2], capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stderr
    scenario = xlsx_to_simulation_design(workbook).select()
    expected = resolve_pressure_inputs(scenario, xlsx_to_well_model(workbook).spec.well_header)
    deck = (tmp_path / "case/model/xlsx_test_default.in").read_text(encoding="utf-8")
    assert f"DATUM_D  {expected.z_fluid_contact:g} m" in deck
    assert f"WGC_D  {expected.z_fluid_contact:g} m" in deck
    assert f"PRESSURE  {expected.p_fluid_contact:g} Bar" in deck


BOX = re.compile(r"^(FIPLEG|PORV)\s+(\S+)(?:\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+))?\s*/\s*$")


def _apply_boxes(text: str, nx: int, ny: int, nz: int):
    """FIPLEG regions and PORV multipliers (k, j, i) that the generated GRDECL boxes produce, in file order."""
    fipleg = np.zeros((nz, ny, nx), dtype=int)
    porv = np.ones((nz, ny, nx))
    for line in text.splitlines():
        match = BOX.match(line.strip())
        if not match:
            continue
        keyword, value, *box = match.groups()
        i0, i1, j0, j1, k0, k1 = (int(v) for v in box) if box[0] else (1, nx, 1, ny, 1, nz)
        window = (slice(k0 - 1, k1), slice(j0 - 1, j1), slice(i0 - 1, i1))
        if keyword == "FIPLEG":
            fipleg[window] = int(value)
        else:
            porv[window] *= float(value)
    return fipleg, porv


@pytest.mark.parametrize("aquifer_layers", [0, 3])
@pytest.mark.parametrize("wall_multiplier", [1.0, 5.0])
def test_pore_volume_and_fipleg_regions_of_the_aquifer(tmp_path, aquifer_layers, wall_multiplier):
    root = Path(__file__).parents[2]
    grdecl = tmp_path / "grid.grdecl"
    grdecl.write_text((root / "test_data/examples/wildcat-pflotran/include/TEMP_GRD.grdecl").read_text(encoding="utf-8"))
    prepare_init_case.parameterize_grdecl(
        grdecl, water_layers=3, overburden_layers=15, reservoir_layers=40, aquifer_layers=aquifer_layers,
        aquifer_permx=250.0, water_overburden_wall_porv_multiplier=1500.0, reservoir_bottom_porv_multiplier=700.0,
        reservoir_wall_porv_multiplier=wall_multiplier,
    )
    text = grdecl.read_text(encoding="utf-8")
    fipleg, porv = _apply_boxes(text, 20, 20, 58)
    edge = np.zeros((20, 20), dtype=bool)
    edge[0, :] = edge[-1, :] = edge[:, 0] = edge[:, -1] = True
    upper, reservoir, bottom = slice(0, 18), slice(18, 58), 57

    # Pore volume: three independent multipliers; nothing else is multiplied.
    assert (porv[upper][:, edge] == 1500).all() and (porv[upper][:, ~edge] == 1).all()
    expected_bottom = np.full((20, 20), 700.0)
    expected_bottom[9, 9] = 1.0  # the cell under the well (and its LGR) is not multiplied
    assert (porv[bottom] == expected_bottom).all()
    assert (porv[18:bottom][:, edge] == wall_multiplier).all() and (porv[18:bottom][:, ~edge] == 1).all()

    # FIPLEG: 7 = water/overburden side walls, 5 = bottom layer; regions 1, 2, 3 keep the rest; region 4 never survives.
    assert (fipleg[upper][:, edge] == 7).all()
    assert (fipleg[0:3][:, ~edge] == 1).all() and (fipleg[3:18][:, ~edge] == 2).all()
    assert (fipleg[bottom] == 5).all()
    assert (fipleg[18:bottom] == 3).all()
    assert 4 not in fipleg
    # Physical aquifer permeability is unchanged: the bottom aquifer_layers keep aquifer_permx.
    assert ("PERMX 250 1 20 1 20 56 58 /" in text) == bool(aquifer_layers)
    # The optional reservoir wall boxes only exist when the multiplier is not 1, and never touch the bottom layer.
    assert ("Optional pore volume multiplier" in text) == (wall_multiplier != 1)


@pytest.mark.parametrize(("nx", "ny", "centre"), [(20, 20, (10, 10)), (21, 6, (11, 3)), (3, 3, (2, 2)), (1, 1, (1, 1))])
def test_the_wells_parent_cell_is_the_only_bottom_cell_without_the_multiplier(tmp_path, nx, ny, centre):
    root = Path(__file__).parents[2]
    grdecl = tmp_path / "grid.grdecl"
    grdecl.write_text((root / "test_data/examples/wildcat-pflotran/include/TEMP_GRD.grdecl").read_text(encoding="utf-8"))
    prepare_init_case.parameterize_grdecl(
        grdecl, water_layers=3, overburden_layers=15, reservoir_layers=40, nx=nx, ny=ny, reservoir_bottom_porv_multiplier=700.0
    )
    fipleg, porv = _apply_boxes(grdecl.read_text(encoding="utf-8"), nx, ny, 58)
    expected = np.full((ny, nx), 700.0)
    expected[centre[1] - 1, centre[0] - 1] = 1.0
    assert (porv[57] == expected).all()
    assert (fipleg[57] == 5).all()  # region 5 still labels the whole bottom layer
    assert (porv[18:57] == 1).all()


def test_single_layer_reservoir_has_no_reservoir_wall_boxes(tmp_path):
    root = Path(__file__).parents[2]
    grdecl = tmp_path / "grid.grdecl"
    grdecl.write_text((root / "test_data/examples/wildcat-pflotran/include/TEMP_GRD.grdecl").read_text(encoding="utf-8"))
    prepare_init_case.parameterize_grdecl(
        grdecl, water_layers=3, overburden_layers=15, reservoir_layers=1, aquifer_layers=0, reservoir_wall_porv_multiplier=5.0
    )
    _, porv = _apply_boxes(grdecl.read_text(encoding="utf-8"), 20, 20, 19)
    expected = np.full((20, 20), 2000.0)
    expected[9, 9] = 1.0
    assert (porv[18] == expected).all()


@pytest.mark.parametrize(
    ("kwargs", "permx", "poro"), [({}, "10000", "1"), ({"free_fluid_permx": 5000.0, "free_fluid_poro": 0.9}, "5000", "0.9")]
)
def test_ocean_layers_use_the_free_fluid_definition(tmp_path, kwargs, permx, poro):
    root = Path(__file__).parents[2]
    grdecl = tmp_path / "grid.grdecl"
    grdecl.write_text((root / "test_data/examples/wildcat-pflotran/include/TEMP_GRD.grdecl").read_text(encoding="utf-8"))
    prepare_init_case.parameterize_grdecl(grdecl, water_layers=3, overburden_layers=15, reservoir_layers=40, **kwargs)
    text = grdecl.read_text(encoding="utf-8")
    assert f"PERMX {permx} 1 20 1 20 1 3 /" in text
    assert f"PORO {poro} 1 20 1 20 1 3 /" in text


def test_standalone_staging_applies_designmatrix_grid_overrides(tmp_path):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "wb.xlsx"
    shutil.copyfile(root / "test_data/examples/wildcat/wildcat_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    excel["GridPolicy"].append(["reservoir_bottom_porv_multiplier", 700])
    sheet = excel["DesignMatrix"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for name, value in (("reservoir_bottom_porv_multiplier", 900), ("reservoir_wall_porv_multiplier", 5)):
        sheet.cell(2, columns[name]).value = value
    excel.save(workbook)

    result = subprocess.run(
        [
            sys.executable, "runscripts/prepare_init_case_from_xlsx.py", "--xlsx", str(workbook),
            "--output-root", str(tmp_path / "case"), "--case-name", "baseline",
        ],
        cwd=root, capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stderr
    grid = next((tmp_path / "case/include").glob("*_GRD.grdecl")).read_text(encoding="utf-8")
    boxes = re.findall(r"^PORV 900 (\d+) (\d+) (\d+) (\d+) (\d+) \5 /$", grid, re.MULTILINE)  # the row beats GridPolicy's 700
    assert len(boxes) == 4 and sum((int(b) - int(a) + 1) * (int(d) - int(c) + 1) for a, b, c, d, _ in boxes) == 399
    assert re.search(r"^PORV 5 1 1 1 20 \d+ \d+ /$", grid, re.MULTILINE)


def test_xlsx_grid_policy_has_free_fluid_defaults(tmp_path):
    workbook = tmp_path / "defaults.xlsx"
    _write_minimal_workbook(workbook)
    policy = xlsx_grid_policy(workbook)
    assert (policy["free_fluid_permx"], policy["free_fluid_poro"]) == (10000.0, 1.0)
    assert (policy["porv_multiplier"], policy["reservoir_wall_porv_multiplier"]) == (2000.0, 1.0)
    assert (policy["water_overburden_wall_porv_multiplier"], policy["reservoir_bottom_porv_multiplier"]) == (2000.0, 2000.0)


def _grid_policy_with(tmp_path, **extra):
    workbook = tmp_path / "policy.xlsx"
    _write_minimal_workbook(workbook)
    excel = load_workbook(workbook)
    for key, value in extra.items():
        excel["GridPolicy"].append([key, value])
    excel.save(workbook)
    return xlsx_grid_policy(workbook)


def test_legacy_porv_multiplier_is_the_default_of_the_two_large_multipliers(tmp_path):
    policy = _grid_policy_with(tmp_path, porv_multiplier=500)
    assert (policy["water_overburden_wall_porv_multiplier"], policy["reservoir_bottom_porv_multiplier"]) == (500, 500)
    assert policy["reservoir_wall_porv_multiplier"] == 1.0
    policy = _grid_policy_with(tmp_path, porv_multiplier=500, reservoir_bottom_porv_multiplier=700)
    assert (policy["water_overburden_wall_porv_multiplier"], policy["reservoir_bottom_porv_multiplier"]) == (500, 700)


def test_workbook_keeps_simulation_assumptions_outside_well_model(tmp_path):
    workbook = tmp_path / "well_input.xlsx"
    _write_minimal_workbook(workbook)

    well_model = xlsx_to_well_model(workbook)
    scenario = xlsx_to_simulation_design(workbook).select()

    assert well_model.spec.subsurface_assumptions is None
    assert scenario.z_fluid_contact == 2400.0
    assert scenario.p_fluid_contact == 210.0


def test_prepare_init_case_from_xlsx_configures_final_run(tmp_path):
    workbook = tmp_path / "well_input.xlsx"
    _write_minimal_workbook(workbook, salinity=0.04)
    output_root = tmp_path / "staged_case"
    repo_root = Path(__file__).parents[2]

    command = [
        sys.executable,
        "runscripts/prepare_init_case_from_xlsx.py",
        "--xlsx",
        str(workbook),
        "--output-root",
        str(output_root),
        "--final-run",
    ]
    subprocess.run(command, check=True, cwd=repo_root, capture_output=True, text=True)

    deck = (output_root / "model" / "xlsx_test_default.in").read_text(encoding="utf-8")
    grdecl = (output_root / "include" / "xlsx_test_default_GRD.grdecl").read_text(encoding="utf-8")
    assert "FINAL_DATE  1 JAN 2125" in deck
    assert "DATUM_D  500 m" in deck
    pressure_table = PressureTable(
        name="expected",
        depth=np.arange(0.0, 510.0, 10.0),
        ground_elevation=105.0,
        ground_temperature=4.0,
        geothermal_gradient=31.0,
        salinity=0.04,
    )
    expected_pressure = pressure_table.get_values_at_depth(500.0)["hydrostatic_pressure"]
    assert f"PRESSURE  {expected_pressure:g} Bar" in deck
    assert "DATUM_D  2400 m" in deck
    assert "PRESSURE  210 Bar" in deck
    assert "WGC_D  2400 m" in deck
    assert "     4    4" in deck
    assert "     105    4" in deck
    assert "     2400    75.145" in deck
    assert deck.count("     2400    75.145") == 2
    assert "\n BRINE  0.04 MASS\n" in deck
    assert deck.count("CONCENTRATION_UNITS MASS") == 2
    assert deck.count("SALTVD\n     4 0.04\n     2400 0.04") == 2
    assert "WELL_DATA INJ_01" not in deck
    assert deck.count("# SCREEN case:") == 1
    assert "external_file ../include/xlsx_test_default_LGR.grdecl /" in grdecl
    assert "TEMP_LGR" not in grdecl


def test_xlsx_grid_policy_requires_keys(tmp_path):
    workbook = tmp_path / "bad_policy.xlsx"
    header = pd.DataFrame({"key": ["unique_wellbore_identifier"], "value": ["NO 00/0-0"]})
    policy = pd.DataFrame({"key": ["top_depth"], "value": [4.0]})
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        header.to_excel(writer, sheet_name="Header", index=False)
        policy.to_excel(writer, sheet_name="GridPolicy", index=False)

    try:
        xlsx_grid_policy(workbook)
    except ValueError as exc:
        assert "missing required keys" in str(exc)
    else:
        raise AssertionError("Expected ValueError for missing GridPolicy keys")


def test_xlsx_grid_policy_applies_legacy_units_and_defaults(tmp_path):
    workbook = tmp_path / "legacy_policy.xlsx"
    _write_minimal_workbook(workbook)

    policy = xlsx_grid_policy(workbook)

    assert policy["depth_unit"] == "m"
    assert policy["permeability_unit"] == "mD"
    assert policy["dx"] == 200.0
    assert policy["reservoir_permx"] == 1000.0


@pytest.mark.parametrize(
    ("key", "value"),
    [
        ("depth_unit", "ft"), ("permeability_unit", "D"), ("target_dz_water", 0), ("reservoir_permx", -1),
        ("free_fluid_permx", 0), ("free_fluid_poro", 0), ("free_fluid_poro", 1.5),
        ("porv_multiplier", 0), ("reservoir_wall_porv_multiplier", 0),
        ("water_overburden_wall_porv_multiplier", 0), ("reservoir_bottom_porv_multiplier", -1),
    ],
)
def test_xlsx_grid_policy_rejects_invalid_values(tmp_path, key, value):
    workbook = tmp_path / "invalid_policy.xlsx"
    _write_minimal_workbook(workbook)
    excel = load_workbook(workbook)
    sheet = excel["GridPolicy"]
    matching_rows = [row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == key]
    if matching_rows:
        sheet.cell(matching_rows[0], 2).value = value
    else:
        sheet.append([key, value])
    excel.save(workbook)

    with pytest.raises(ValueError, match="invalid GridPolicy"):
        xlsx_grid_policy(workbook)


def test_xlsx_grid_policy_converts_header_elevation_from_feet(tmp_path):
    workbook = tmp_path / "feet_header.xlsx"
    _write_minimal_workbook(workbook)
    excel = load_workbook(workbook)
    sheet = excel["Header"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "ground_elevation":
            sheet.cell(row, 2).value = 1000.0
        elif sheet.cell(row, 1).value == "ground_elevation_unit":
            sheet.cell(row, 2).value = "ft"
    excel.save(workbook)

    assert xlsx_grid_policy(workbook)["water_depth"] == pytest.approx(304.8)


def test_case_name_defaults_to_default_when_omitted(tmp_path):
    """Verify that --case-name defaults to 'default' when omitted."""
    workbook = tmp_path / "well_input.xlsx"
    _write_minimal_workbook(workbook)

    output_root = tmp_path / "staged_case"
    repo_root = Path(__file__).parents[2]

    command = [
        sys.executable,
        "runscripts/prepare_init_case_from_xlsx.py",
        "--xlsx",
        str(workbook),
        "--output-root",
        str(output_root),
    ]
    subprocess.run(command, check=True, cwd=repo_root, capture_output=True, text=True)

    # Should succeed without specifying --case-name
    assert (output_root / "model" / "xlsx_test_default.in").exists()


def test_case_name_selection_with_multi_scenario(tmp_path):
    """Verify that --case-name selects the correct scenario from multi-scenario workbook."""
    workbook = tmp_path / "multi_scenario.xlsx"
    metadata = pd.DataFrame({"key": ["namespace", "name", "author"], "value": ["screen", "xlsx-test", "pytest"]})
    header = pd.DataFrame(
        {
            "key": [
                "unique_wellbore_identifier",
                "depth_reference_rkb",
                "depth_reference_rkb_unit",
                "ground_elevation",
                "ground_elevation_unit",
                "total_depth_rkb",
                "total_depth_rkb_unit",
            ],
            "value": ["NO 00/0-0", 27, "m", 105, "m", 3997, "m"],
        }
    )
    grid_policy = pd.DataFrame(
        {
            "key": [
                "top_depth",
                "water_depth",
                "reservoir_top",
                "bottom_depth",
                "target_dz_water",
                "target_dz_overburden",
                "target_dz_reservoir",
                "cells_per_layer",
            ],
            "value": [4.0, 104.0, 1004.0, 1504.0, 50.0, 60.0, 10.0, 400],
        }
    )

    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        metadata.to_excel(writer, sheet_name="Metadata", index=False)
        header.to_excel(writer, sheet_name="Header", index=False)
        grid_policy.to_excel(writer, sheet_name="GridPolicy", index=False)
        pd.DataFrame(
            {
                "case_name": ["default", "hot_case"],
                "temperature_gradient": [31.0, 40.0],
                "ground_temperature": [4.0, 5.0],
                "z_fluid_contact": [2400.0, 2350.0],
                "p_fluid_contact": [210.0, 220.0],
                "overburden_datum_depth": [500.0, 500.0],
                "z_resrv": [1400.0, 1400.0],
                "p_resrv": [250.0, 260.0],
            }
        ).to_excel(writer, sheet_name="SubsurfaceAssumptions", index=False)

    design = xlsx_to_simulation_design(workbook)
    default_scenario = design.select("default")
    hot_scenario = design.select("hot_case")

    assert default_scenario.temperature_gradient == 31.0
    assert hot_scenario.temperature_gradient == 40.0


def test_case_name_raises_error_for_unknown_case(tmp_path):
    """Verify that selecting an unknown case name raises an error."""
    workbook = tmp_path / "well_input.xlsx"
    _write_minimal_workbook(workbook)

    design = xlsx_to_simulation_design(workbook)
    try:
        design.select("nonexistent_case")
    except ValueError as exc:
        assert "unknown simulation case" in str(exc)
    else:
        raise AssertionError("Expected ValueError for unknown case name")
