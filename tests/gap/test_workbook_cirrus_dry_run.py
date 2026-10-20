from pathlib import Path
import json
import importlib.util
import re
import shutil
import subprocess
import sys
import types
from types import SimpleNamespace

import numpy as np
import pytest
from openpyxl import load_workbook

from src.GaP.libs.models.simulation_scenario import SimulationScenario
from src.WellClass.libs.utils.xlsx_parser import xlsx_to_simulation_design, xlsx_to_well_model

WORKFLOW_PATH = Path(__file__).parents[2] / "runscripts/run_workbook_to_cirrus_lgr.py"
sys.path.insert(0, str(WORKFLOW_PATH.parent))
WORKFLOW_SPEC = importlib.util.spec_from_file_location("run_workbook_to_cirrus_lgr", WORKFLOW_PATH)
workflow = importlib.util.module_from_spec(WORKFLOW_SPEC)
assert WORKFLOW_SPEC.loader is not None
WORKFLOW_SPEC.loader.exec_module(workflow)


def _fake_cirrus(tmp_path: Path, fixture_prefix: Path) -> Path:
    runner = tmp_path / "fake-cirrus"
    runner.write_text(
        "#!/bin/sh\n"
        f'cp "{fixture_prefix}.EGRID" "${{1%.in}}.EGRID"\n'
        f'cp "{fixture_prefix}.INIT" "${{1%.in}}.INIT"\n'
        'echo "fake CIRRUS completed"\n',
        encoding="utf-8",
    )
    runner.chmod(0o755)
    return runner


def test_workbook_wrapper_runs_full_dry_run_without_cirrus(tmp_path):
    root = Path(__file__).parents[2]
    output_root = tmp_path / "case"
    fixture_prefix = root / "test_data/examples/wildcat/model/TEMP-0"
    runner = _fake_cirrus(tmp_path, fixture_prefix)

    subprocess.run(
        [
            sys.executable,
            "runscripts/run_workbook_to_cirrus_lgr.py",
            "--xlsx",
            "test_data/examples/wildcat/wildcat_workbook.xlsx",
            "--output-root",
            str(output_root),
            "--sim-command",
            f"{runner} {{deck}}",
            "--run-final",
            "--simulation-years",
            "0",
            "--case-name",
            "baseline",
            "--plot",
        ],
        cwd=root,
        check=True,
        capture_output=True,
        text=True,
    )

    assert (output_root / "well_input.json").exists()
    scenario = json.loads((output_root / "scenario.json").read_text(encoding="utf-8"))
    grid_policy = json.loads((output_root / "grid_policy.json").read_text(encoding="utf-8"))
    assert scenario["case_name"] == "baseline"
    assert scenario["temperature_gradient"] == 31.0
    assert scenario["case_metadata"] == {
        "case_stem": "wildcat_baseline",
        "case_index": 1,
        "case_name": "baseline",
        "well_label": "wildcat",
        "workbook": "test_data/examples/wildcat/wildcat_workbook.xlsx",
        "template_root": "test_data/examples/wildcat-pflotran",
        "template_files": {"deck": "model/TEMP-0.in", "grid": "include/TEMP_GRD.grdecl"},
        "case_files": {
            "deck": "model/wildcat_baseline.in",
            "grid": "include/wildcat_baseline_GRD.grdecl",
            "lgr": "include/wildcat_baseline_LGR.grdecl",
        },
    }
    assert grid_policy["top_depth"] == 4.0
    assert (output_root / "qc_plot.png").stat().st_size > 0
    assert (output_root / "model/wildcat_baseline.EGRID").exists()
    assert (output_root / "model/wildcat_baseline.INIT").exists()
    assert (output_root / "include/wildcat_baseline_LGR.grdecl").exists()
    assert not list(output_root.rglob("TEMP*"))
    assert (output_root / "logs/wildcat_baseline_initialization.log").read_text(encoding="utf-8") == "fake CIRRUS completed\n"
    assert (output_root / "logs/wildcat_baseline_final.log").read_text(encoding="utf-8") == "fake CIRRUS completed\n"
    deck = (output_root / "model/wildcat_baseline.in").read_text(encoding="utf-8")
    assert "FINAL_DATE  1 JAN 2025" in deck
    assert f"\n BRINE  {scenario['salinity']:g} MASS\n" in deck
    assert deck.count("# SCREEN case: wildcat_baseline | scenario 'baseline' (DesignMatrix row 1)") == 1
    assert "external_file ../include/wildcat_baseline_LGR.grdecl /" in (
        output_root / "include/wildcat_baseline_GRD.grdecl"
    ).read_text(encoding="utf-8")


def test_design_matrix_permeability_overrides_reach_case_outputs(tmp_path):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "design_matrix.xlsx"
    shutil.copyfile(root / "test_data/examples/wildcat/wildcat_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    sheet = excel["DesignMatrix"]
    columns = {sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)}
    for field, value in {
        "reservoir_permx": 750.0,
        "overburden_permx": 0.002,
        "cb_perm": 0.2,
        "barrier_perm": 0.3,
    }.items():
        sheet.cell(2, columns[field]).value = value
    excel.save(workbook)

    output_root = tmp_path / "overridden_case"
    fixture_prefix = root / "test_data/examples/wildcat/model/TEMP-0"
    runner = _fake_cirrus(tmp_path, fixture_prefix)
    subprocess.run(
        [
            sys.executable,
            "runscripts/run_workbook_to_cirrus_lgr.py",
            "--xlsx",
            str(workbook),
            "--output-root",
            str(output_root),
            "--sim-command",
            f"{runner} {{deck}}",
            "--case-name",
            "baseline",
        ],
        cwd=root,
        check=True,
        capture_output=True,
        text=True,
    )

    scenario = json.loads((output_root / "scenario.json").read_text(encoding="utf-8"))
    policy = json.loads((output_root / "grid_policy.json").read_text(encoding="utf-8"))
    coarse_grid = (output_root / "include/wildcat_baseline_GRD.grdecl").read_text(encoding="utf-8")
    lgr = (output_root / "include/wildcat_baseline_LGR.grdecl").read_text(encoding="utf-8")
    assert scenario["effective_permeability_mD"] == {
        "reservoir_permx": 750.0,
        "overburden_permx": 0.002,
        "oh_perm": 10000.0,
        "cb_perm": 0.2,
        "cb_perm_interval_overrides": {},
        "barrier_perm": 0.3,
        "barrier_perm_interval_overrides": {},
    }
    assert policy["reservoir_permx"] == 750.0
    assert policy["overburden_permx"] == 0.002
    assert "PERMX 750" in coarse_grid
    assert "PERMX  0.2" in lgr
    assert "PERMX  0.3" in lgr
    assert "CARFIN\nTEMP_LGR 10 10 10 10 1 60 22 22 150 /" in lgr


def test_unknown_interval_target_is_rejected_before_cirrus_runs():
    root = Path(__file__).parents[2]
    model = xlsx_to_well_model(root / "test_data/examples/smeaheia/smeaheia_workbook.xlsx")
    scenario = SimulationScenario(barrier_perm=0.2, barrier_perm_interval="missing")

    with pytest.raises(ValueError, match="unknown plug interval 'missing'"):
        workflow.validate_scenario_interval_targets(model, scenario)


def test_workbook_interval_permeability_overrides_are_isolated(tmp_path):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "targeted.xlsx"
    shutil.copyfile(root / "test_data/examples/smeaheia/smeaheia_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    sheet = excel["DesignMatrix"]
    columns = {sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)}
    for field, value in {
        "cb_perm": 0.25,
        "cb_perm_interval": "Cement 9 5/8 in",
        "barrier_perm": 0.75,
        "barrier_perm_interval": "cplug9",
    }.items():
        sheet.cell(2, columns[field]).value = value
    excel.save(workbook)

    fixture_prefix = root / "test_data/examples/smeaheia/model/TEMP-0"
    runner = _fake_cirrus(tmp_path, fixture_prefix)
    output_root = tmp_path / "targeted"
    subprocess.run(
        [
            sys.executable,
            "runscripts/run_workbook_to_cirrus_lgr.py",
            "--xlsx",
            str(workbook),
            "--output-root",
            str(output_root),
            "--sim-command",
            f"{runner} {{deck}}",
            "--case-name",
            "baseline",
        ],
        cwd=root,
        check=True,
        capture_output=True,
        text=True,
    )

    lgr = (output_root / "include/smeaheia_baseline_LGR.grdecl").read_text(encoding="utf-8")
    assert "PERMX  0.25  9  14  9  9  49  91  /" in lgr
    assert "PERMX  0.75  9  14  9  14  14  35  /" in lgr
    scenario = json.loads((output_root / "scenario.json").read_text(encoding="utf-8"))
    assert scenario["effective_permeability_mD"]["cb_perm"] == 0.05
    assert scenario["effective_permeability_mD"]["cb_perm_interval_overrides"] == {"Cement 9 5/8 in": 0.25}
    assert scenario["effective_permeability_mD"]["barrier_perm"] == 0.05
    assert scenario["effective_permeability_mD"]["barrier_perm_interval_overrides"] == {"cplug9": 0.75}


def _hole_workbook(tmp_path, casing, **extra):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "hole.xlsx"
    shutil.copyfile(root / "test_data/examples/wildcat/wildcat_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    sheet = excel["DesignMatrix"]
    columns = {sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)}
    for field, value in {
        "casing_hole_casing": casing, "casing_hole_depth_mTVDMSL": 200, "casing_hole_diameter_m": 0.1, **extra,
    }.items():
        if field not in columns:
            columns[field] = sheet.max_column + 1
            sheet.cell(1, columns[field]).value = field
        sheet.cell(2, columns[field]).value = value
    excel.save(workbook)
    return workbook


@pytest.mark.parametrize("casing", ["missing", "Cement 20 in"])
def test_invalid_hole_target_is_rejected_before_initialization(tmp_path, monkeypatch, casing):
    workbook = _hole_workbook(tmp_path, casing)
    monkeypatch.setattr(workflow, "stage_case", lambda *_: pytest.fail("invalid hole must not stage or run CIRRUS"))
    args = SimpleNamespace(
        xlsx=workbook, case_name="baseline", sim_command="never {deck}",
        queue_poll_interval=15, queue_timeout=60,
    )
    with pytest.raises(ValueError, match="exactly one casing"):
        workflow.run_workflow(args)


def test_workbook_hole_reaches_lgr_and_resolved_metadata(tmp_path):
    root = Path(__file__).parents[2]
    workbook = _hole_workbook(tmp_path, "Casing 20 in")
    runner = _fake_cirrus(tmp_path, root / "test_data/examples/wildcat/model/TEMP-0")
    output_root = tmp_path / "hole_case"
    subprocess.run(
        [
            sys.executable, "runscripts/run_workbook_to_cirrus_lgr.py",
            "--xlsx", str(workbook), "--output-root", str(output_root),
            "--sim-command", f"{runner} {{deck}}", "--case-name", "baseline",
        ],
        cwd=root, check=True, capture_output=True, text=True,
    )
    scenario = json.loads((output_root / "scenario.json").read_text())
    report = scenario["resolved_casing_hole"]
    lgr_path = output_root / "include/wildcat_baseline_LGR.grdecl"
    assert report == json.loads(lgr_path.with_suffix(".casing_hole.json").read_text())
    assert report["experimental"] is True
    assert report["input"] == {"casing": "Casing 20 in", "depth_mTVDMSL": 200, "diameter_m": 0.1}
    assert report["side"] == "+X"
    assert report["allocated_area_m2"] == pytest.approx(np.pi * 0.05**2, rel=1e-12)
    assert report["allocation"] == "circular area weighted by lateral span overlap"
    assert [(face["i"], face["j"], face["k"]) for face in report["faces"]] == [(15, 11, 17), (15, 12, 17)]
    assert [face["lateral_overlap_m"] for face in report["faces"]] == pytest.approx([0.05, 0.05])
    assert [face["opening_area_m2"] for face in report["faces"]] == pytest.approx([np.pi * 0.05**2 / 2] * 2)
    for face in report["faces"]:
        j = face["j"]
        assert f"MULTX {face['multiplier']:.17g} 15 15 {j} {j} 17 17 /" in lgr_path.read_text()


def test_workbook_report_only_hole_keeps_fiphol_and_stays_sealed(tmp_path):
    root = Path(__file__).parents[2]
    runner = _fake_cirrus(tmp_path, root / "test_data/examples/wildcat/model/TEMP-0")
    lgrs = {}
    for name, extra in {"open": {}, "sealed": {"casing_hole_open": False}}.items():
        (tmp_path / name).mkdir()
        workbook = _hole_workbook(tmp_path / name, "Casing 20 in", **extra)
        output_root = tmp_path / f"{name}_case"
        subprocess.run(
            [
                sys.executable, "runscripts/run_workbook_to_cirrus_lgr.py",
                "--xlsx", str(workbook), "--output-root", str(output_root),
                "--sim-command", f"{runner} {{deck}}", "--case-name", "baseline",
            ],
            cwd=root, check=True, capture_output=True, text=True,
        )
        lgrs[name] = (output_root / "include/wildcat_baseline_LGR.grdecl").read_text()
        report = json.loads((output_root / "scenario.json").read_text())["resolved_casing_hole"]
        assert report["opened"] is (name == "open")
        assert report["allocated_area_m2"] == (pytest.approx(np.pi * 0.05**2) if name == "open" else 0.0)
    fiphol = [line for line in lgrs["open"].splitlines() if line.startswith("FIPHOL")]
    assert len(fiphol) == 5
    assert fiphol == [line for line in lgrs["sealed"].splitlines() if line.startswith("FIPHOL")]
    assert "SCREEN experimental casing hole" in lgrs["open"]
    assert "SCREEN experimental casing hole" not in lgrs["sealed"]


def _run_wrapper(tmp_path, workbook, output_root, *extra):
    root = Path(__file__).parents[2]
    runner = _fake_cirrus(tmp_path, root / "test_data/examples/wildcat/model/TEMP-0")
    return subprocess.run(
        [
            sys.executable, "runscripts/run_workbook_to_cirrus_lgr.py",
            "--xlsx", str(workbook), "--output-root", str(output_root),
            "--sim-command", f"{runner} {{deck}}", "--case-name", "baseline", *extra,
        ],
        cwd=root, capture_output=True, text=True,
    )


def test_gridpolicy_free_fluid_values_reach_ocean_and_wellbore_cells(tmp_path):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "free_fluid.xlsx"
    shutil.copyfile(root / "test_data/examples/wildcat/wildcat_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    excel["GridPolicy"].append(["free_fluid_permx", 5000])
    excel["GridPolicy"].append(["free_fluid_poro", 0.9])
    excel.save(workbook)

    result = _run_wrapper(tmp_path, workbook, tmp_path / "case")
    assert result.returncode == 0, result.stderr
    grid = (tmp_path / "case/include/wildcat_baseline_GRD.grdecl").read_text(encoding="utf-8")
    lgr = (tmp_path / "case/include/wildcat_baseline_LGR.grdecl").read_text(encoding="utf-8")
    assert re.search(r"^PERMX 5000 1 20 1 20 1 \d+ /$", grid, re.MULTILINE)
    assert re.search(r"^PORO 0.9 1 20 1 20 1 \d+ /$", grid, re.MULTILINE)
    assert "PERMX  5000.0  9  14  9  14  83  150  /" in lgr
    assert "PORO  0.9  9  14  9  14  83  150  /" in lgr
    policy = json.loads((tmp_path / "case/grid_policy.json").read_text(encoding="utf-8"))
    assert (policy["free_fluid_permx"], policy["free_fluid_poro"]) == (5000, 0.9)

    conflict = _run_wrapper(tmp_path, workbook, tmp_path / "conflict", "--oh-perm", "7000")
    assert conflict.returncode != 0
    assert "conflicts with the free-fluid permeability 5000 mD" in conflict.stderr


def test_porv_multipliers_come_from_gridpolicy_and_designmatrix_overrides_win(tmp_path):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "porv.xlsx"
    shutil.copyfile(root / "test_data/examples/wildcat/wildcat_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    for key, value in (("water_overburden_wall_porv_multiplier", 1500), ("reservoir_bottom_porv_multiplier", 700), ("reservoir_wall_porv_multiplier", 3)):
        excel["GridPolicy"].append([key, value])
    sheet = excel["DesignMatrix"]
    sheet.cell(1, sheet.max_column + 1).value = "reservoir_bottom_porv_multiplier"
    sheet.cell(2, sheet.max_column).value = 900
    excel.save(workbook)

    result = _run_wrapper(tmp_path, workbook, tmp_path / "case")
    assert result.returncode == 0, result.stderr
    grid = (tmp_path / "case/include/wildcat_baseline_GRD.grdecl").read_text(encoding="utf-8")
    assert re.search(r"^PORV 1500 1 1 1 20 1 \d+ /$", grid, re.MULTILINE)
    assert re.search(r"^PORV 900 1 20 1 20 (\d+) \1 /$", grid, re.MULTILINE)  # the row beats GridPolicy's 700
    assert re.search(r"^PORV 3 1 1 1 20 \d+ \d+ /$", grid, re.MULTILINE)
    policy = json.loads((tmp_path / "case/grid_policy.json").read_text(encoding="utf-8"))
    assert (
        policy["water_overburden_wall_porv_multiplier"], policy["reservoir_bottom_porv_multiplier"],
        policy["reservoir_wall_porv_multiplier"],
    ) == (1500, 900, 3)
    scenario = json.loads((tmp_path / "case/scenario.json").read_text(encoding="utf-8"))
    assert scenario["reservoir_bottom_porv_multiplier"] == 900 and scenario["reservoir_wall_porv_multiplier"] is None


@pytest.mark.parametrize(
    "field", ["water_overburden_wall_porv_multiplier", "reservoir_bottom_porv_multiplier", "reservoir_wall_porv_multiplier"]
)
def test_scenario_rejects_non_positive_porv_overrides(field):
    with pytest.raises(ValueError):
        SimulationScenario(**{field: 0})
    assert getattr(SimulationScenario(), field) is None


def test_designmatrix_free_fluid_override_beats_gridpolicy_for_that_case(tmp_path):
    root = Path(__file__).parents[2]
    workbook = tmp_path / "free_fluid_row.xlsx"
    shutil.copyfile(root / "test_data/examples/wildcat/wildcat_workbook.xlsx", workbook)
    excel = load_workbook(workbook)
    excel["GridPolicy"].append(["free_fluid_permx", 5000])
    sheet = excel["DesignMatrix"]
    for column, value in ((sheet.max_column + 1, 2000), (sheet.max_column + 2, 0.8)):
        sheet.cell(1, column).value = "free_fluid_permx" if value == 2000 else "free_fluid_poro"
        sheet.cell(2, column).value = value
    excel.save(workbook)

    result = _run_wrapper(tmp_path, workbook, tmp_path / "case")
    assert result.returncode == 0, result.stderr
    grid = (tmp_path / "case/include/wildcat_baseline_GRD.grdecl").read_text(encoding="utf-8")
    lgr = (tmp_path / "case/include/wildcat_baseline_LGR.grdecl").read_text(encoding="utf-8")
    assert re.search(r"^PERMX 2000 1 20 1 20 1 \d+ /$", grid, re.MULTILINE)
    assert re.search(r"^PORO 0.8 1 20 1 20 1 \d+ /$", grid, re.MULTILINE)
    assert "PERMX  2000.0  9  14  9  14  83  150  /" in lgr
    assert "PORO  0.8  9  14  9  14  83  150  /" in lgr
    policy = json.loads((tmp_path / "case/grid_policy.json").read_text(encoding="utf-8"))
    scenario = json.loads((tmp_path / "case/scenario.json").read_text(encoding="utf-8"))
    assert (policy["free_fluid_permx"], policy["free_fluid_poro"]) == (2000, 0.8)
    assert (scenario["free_fluid_permx"], scenario["free_fluid_poro"]) == (2000, 0.8)
    assert scenario["effective_permeability_mD"]["oh_perm"] == 2000


@pytest.mark.parametrize("fields", [{"free_fluid_permx": 0}, {"free_fluid_poro": 0}, {"free_fluid_poro": 1.5}])
def test_scenario_rejects_invalid_free_fluid_overrides(fields):
    with pytest.raises(ValueError):
        SimulationScenario(**fields)
    assert SimulationScenario().free_fluid_permx is None


def test_older_workbook_without_hole_columns_remains_valid(tmp_path):
    workbook = _hole_workbook(tmp_path, "Casing 20 in")
    excel = load_workbook(workbook)
    sheet = excel["DesignMatrix"]
    for column in range(sheet.max_column, 0, -1):
        if str(sheet.cell(1, column).value).startswith("casing_hole_"):
            sheet.delete_cols(column)
    excel.save(workbook)
    assert all(scenario.casing_hole is None for scenario in xlsx_to_simulation_design(workbook).scenarios)


def test_qc_plot_uses_workbook_fluid_contact_pressure(tmp_path, monkeypatch):
    """Verify that the WellClass QC curve uses the scenario pressure datum."""
    captured = {}

    class FakePressureScenario:
        def display_curves(self):
            return {
                "brine_pressure": np.array([100.0, 110.0]),
                "brine_depth": np.array([0.0, 10.0]),
                "fluid_pressure": np.array([200.0, 210.0]),
                "fluid_depth": np.array([0.0, 10.0]),
            }

    class FakePressure:
        def __init__(self, **kwargs):
            captured["pressure_options"] = kwargs
            self.table = SimpleNamespace(
                depth=np.array([0.0, 10.0]),
                hydrostatic_pressure=np.array([100.0, 105.0]),
                min_horizontal_stress=np.array([90.0, 95.0]),
            )

        def add_scenario(self, name, **kwargs):
            captured.update(name=name, **kwargs)
            return FakePressureScenario()

    monkeypatch.setattr(workflow, "Pressure", FakePressure)
    monkeypatch.setattr(workflow, "WellProcessed", SimpleNamespace(from_pydantic=lambda model: SimpleNamespace(header={})))
    import matplotlib.pyplot as plt

    def fake_plot_one_page(well, pressure, scenario_name, **kwargs):
        captured.update(one_page_scenario=scenario_name, one_page_options=kwargs)
        return plt.figure(), None

    monkeypatch.setattr("src.WellClass.libs.plotting.plot_one_page", fake_plot_one_page)

    scenario = SimpleNamespace(
        case_name="pressure_case",
        casing_hole="hole",
        z_fluid_contact=2400.0,
        p_fluid_contact=245.0,
        z_resrv=1400.0,
        p_resrv=250.0,
        ground_temperature=4.0,
        temperature_gradient=31.0,
        fluid_type="co2",
        salinity=0.032,
    )

    workflow.save_qc_plot(SimpleNamespace(), scenario, tmp_path / "qc_plot.png")

    assert captured["z_fluid_datum"] == 2400.0
    assert captured["p_fluid_datum"] == 245.0
    assert captured["pressure_options"]["salinity"] == 0.032
    assert captured["one_page_scenario"] == captured["name"] == "workbook"
    assert captured["one_page_options"]["casing_hole"] == "hole"
    assert captured["one_page_options"]["title"] == "pressure_case WellClass QC"
